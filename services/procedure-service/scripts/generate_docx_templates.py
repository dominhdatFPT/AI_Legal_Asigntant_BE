import os
import io
import sys
from pathlib import Path
import pymupdf
import docx
from docx.shared import Inches

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = Path(r"D:\Legal Assistant AI\BE\services\procedure-service")
BIZ_DIR = BASE_DIR / "Template các thủ tục kinh doanh"
LAND_DIR = BASE_DIR / "Template các thủ tục đất đai"

AUTHOR_NAME = "Cổng Dịch vụ công Quốc gia"

def convert_pdf_to_docx(pdf_path: Path, docx_path: Path):
    print(f"Generating DOCX from PDF: {pdf_path.name} -> {docx_path.name}")
    doc_pdf = pymupdf.open(pdf_path)
    doc_word = docx.Document()
    
    # Configure metadata
    doc_word.core_properties.author = AUTHOR_NAME
    doc_word.core_properties.last_modified_by = AUTHOR_NAME
    doc_word.core_properties.title = docx_path.stem
    
    # Configure A4 page
    section = doc_word.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.2)
    section.bottom_margin = Inches(0.2)
    section.left_margin = Inches(0.2)
    section.right_margin = Inches(0.2)

    for i, page in enumerate(doc_pdf):
        # 150 DPI provides sharp text while keeping file size reasonable
        pix = page.get_pixmap(dpi=150)
        img_data = pix.tobytes("png")
        if i > 0:
            doc_word.add_page_break()
        doc_word.add_picture(io.BytesIO(img_data), width=Inches(7.87))
        
    doc_word.save(docx_path)
    print(f"  Created {docx_path.name} ({docx_path.stat().st_size:,} bytes, Author: '{AUTHOR_NAME}')")

def convert_png_to_docx(png_path: Path, docx_path: Path):
    print(f"Generating DOCX from PNG: {png_path.name} -> {docx_path.name}")
    doc_word = docx.Document()
    
    # Configure metadata
    doc_word.core_properties.author = AUTHOR_NAME
    doc_word.core_properties.last_modified_by = AUTHOR_NAME
    doc_word.core_properties.title = docx_path.stem
    
    section = doc_word.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.2)
    section.bottom_margin = Inches(0.2)
    section.left_margin = Inches(0.2)
    section.right_margin = Inches(0.2)
    
    doc_word.add_picture(str(png_path), width=Inches(7.87))
    doc_word.save(docx_path)
    print(f"  Created {docx_path.name} ({docx_path.stat().st_size:,} bytes, Author: '{AUTHOR_NAME}')")

def update_text_docx_authors():
    print("\n=== UPDATING METADATA FOR TEXT-BASED DOCX FILES ===")
    text_docx_files = [
        LAND_DIR / "Thủ tục cung cấp thông tin, dữ liệu đất đai - Phiếu yêu cầu.docx",
        LAND_DIR / "Thủ tục cấp lại Giấy chứng nhận do bị mất - Đơn đề nghị cấp lại.docx",
        LAND_DIR / "Thủ tục đính chính Giấy chứng nhận đã cấp - Đơn đề nghị đính chính.docx",
        BIZ_DIR / "Thủ tục thay đổi nội dung đăng ký Hộ kinh doanh - Thông báo.docx",
        BIZ_DIR / "Thủ tục thuê nhà - mặt bằng làm Trụ sở chính doanh nghiệp.docx",
        BIZ_DIR / "Thủ tục đăng ký thành lập Hộ kinh doanh - Giấy đề nghị.docx",
    ]
    for p in text_docx_files:
        if p.exists():
            doc = docx.Document(p)
            doc.core_properties.author = AUTHOR_NAME
            doc.core_properties.last_modified_by = AUTHOR_NAME
            doc.core_properties.title = p.stem
            doc.save(p)
            print(f"  [METADATA UPDATED] {p.name} -> Author: '{AUTHOR_NAME}'")

def main():
    targets = [
        # Business PDFs
        (BIZ_DIR, "Thủ tục đăng ký thành lập Công ty Cổ phần.pdf"),
        (BIZ_DIR, "Thủ tục đăng ký thành lập Công ty Hợp danh.pdf"),
        (BIZ_DIR, "Thủ tục đăng ký thành lập Công ty Trách nhiệm hữu hạn.pdf"),
        (BIZ_DIR, "Thủ tục đăng ký thành lập Doanh nghiệp tư nhân.pdf"),
        # Land PDFs & PNG
        (LAND_DIR, "Thủ tục chuyển mục đích sử dụng đất - Đơn đề nghị chuyển mục đích sử dụng đất.pdf"),
        (LAND_DIR, "Thủ tục tách thửa hoặc hợp thửa đất.pdf"),
        (LAND_DIR, "Thủ tục đăng ký biến động quyền sử dụng đất do chuyển nhượng.pdf"),
        (LAND_DIR, "Thủ tục đăng ký đất đai, cấp Giấy chứng nhận lần đầu.pdf"),
        (LAND_DIR, "Thủ tục giao đất, cho thuê đất - Đơn xin giao đất.png"),
    ]
    
    print("=== REGENERATING DOCX TEMPLATES WITH CLEAN SOURCE ===")
    for folder, fname in targets:
        src = folder / fname
        if not src.exists():
            print(f"Warning: File not found {src}")
            continue
        base_name = src.stem
        dest = folder / f"{base_name}.docx"
        if src.suffix.lower() == ".pdf":
            convert_pdf_to_docx(src, dest)
        elif src.suffix.lower() == ".png":
            convert_png_to_docx(src, dest)
            
    update_text_docx_authors()

if __name__ == "__main__":
    main()
