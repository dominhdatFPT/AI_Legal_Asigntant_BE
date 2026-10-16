"""C4 plain-text -> native DOCX -> PDF. No template execution or external input files."""
from io import BytesIO
from pathlib import Path
import re
import shutil
import subprocess
from tempfile import TemporaryDirectory
import unicodedata

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader

from app.core.errors import ContractError


def render_error(code="RENDER_FAILED"):
    return ContractError(code, "Không thể xuất bản này; kiểm tra mẫu và dữ liệu rồi thử lại.", 422)


def plain(value):
    if value is None or value == "":
        return "________"
    if isinstance(value, bool):
        return "Có" if value else "Không"
    return str(value)


class ContractRenderer:
    def ready(self, format):
        # Both exports run PDF QA; do not return an unverified DOCX.
        if not shutil.which("libreoffice"):
            raise ContractError("RENDERER_UNAVAILABLE", "Renderer LibreOffice chưa sẵn sàng; chạy image Docker Contract C7.")

    def docx(self, content, values, preview=False):
        # One substitution pass, no Jinja/eval, values containing {{...}} stay literal.
        text = re.sub(r"\{\{([a-z][a-z0-9_]*)\}\}", lambda m: plain(values.get(m[1])), content)
        if len(text) > 600000 or any(unicodedata.category(c).startswith("C") and c not in "\n\r\t" for c in text):
            raise render_error("INVALID_INPUT")
        if any(len(word) > 150 for word in text.split()):
            raise render_error("TEXT_OVERFLOW")
        document = Document()
        # python-docx's bundled Word template carries a blue paragraph rule.
        # Remove inherited borders rather than merely changing run text/color.
        for border in document.styles.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        for border in document.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        section = document.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(20)
        section.left_margin = section.right_margin = Mm(25)
        style = document.styles["Normal"]
        style.font.name, style.font.size, style.font.color.rgb = "Liberation Serif", Pt(13), RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.widow_control = True
        if preview:
            header = section.header.paragraphs[0]
            header.text = "BẢN XEM TRƯỚC - CHƯA HOÀN THIỆN"
            header.alignment = WD_ALIGN_PARAGRAPH.CENTER
            header.runs[0].bold = True
        for index, line in enumerate(text.replace("\r\n", "\n").replace("\r", "\n").split("\n")):
            paragraph = document.add_paragraph(line)
            if re.match(r"^Điều\s+[0-9]+(?:[ .:]|$)",line) and len(line)<=150:
                paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs:
                    run.bold = True
            if index == 0 and len(line) <= 120:
                paragraph.style = document.styles["Title"]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs:
                    run.font.name, run.font.size, run.font.color.rgb = "Liberation Serif", Pt(16), RGBColor(0,0,0)
                    run.bold = True
        properties = document.core_properties
        properties.author = properties.last_modified_by = properties.title = properties.subject = ""
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def render(self, content, values, format, preview=False):
        self.ready(format)
        docx = self.docx(content, values, preview)
        with TemporaryDirectory(prefix="contract-render-") as temp:
            folder = Path(temp)
            source = folder / "contract.docx"
            source.write_bytes(docx)
            try:
                subprocess.run([shutil.which("libreoffice"), "--headless", "--nologo", "--nodefault", "--norestore",
                    "-env:UserInstallation=" + (folder / "profile").as_uri(), "--convert-to", "pdf:writer_pdf_Export",
                    "--outdir", str(folder), str(source)], check=True, timeout=60,
                    stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                    env={"PATH":"/usr/local/bin:/usr/bin:/bin", "HOME":temp, "TMPDIR":temp, "LANG":"C.UTF-8"})
                pdf = (folder / "contract.pdf").read_bytes()
                self.validate_pdf(pdf)
            except subprocess.TimeoutExpired:
                raise render_error("RENDER_TIMEOUT") from None
            except (OSError, subprocess.CalledProcessError):
                raise render_error() from None
        return docx if format == "docx" else pdf

    def validate_pdf(self, data):
        try:
            if not data.startswith(b"%PDF-") or not 1 <= len(data) <= 52428800:
                raise ValueError()
            reader = PdfReader(BytesIO(data), strict=True)
            if reader.is_encrypted or not 1 <= len(reader.pages) <= 100:
                raise ValueError()
            for page in reader.pages:
                if not (page.extract_text() or "").strip():
                    raise ValueError()
        except Exception:
            raise render_error("INVALID_RENDER_OUTPUT") from None
